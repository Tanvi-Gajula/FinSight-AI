"""Transparent educational FCFF DCF.

This deliberately avoids automatic three-statement forecasting.  Users control
assumptions; Python performs the arithmetic; outputs remain in source units.
"""
from dataclasses import dataclass, replace
import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

@dataclass
class Assumptions:
    growth:float=.06; wacc:float=.11; terminal_growth:float=.04; n:int=5
    shares:float|None=None; market_price:float|None=None

def _latest(df,label):
    if df is None or df.empty or label not in df.index:return None
    for c in sorted(df.columns,key=lambda x:pd.Timestamp(x),reverse=True):
        v=df.loc[label,c]
        if not isinstance(v,pd.Series) and pd.notna(v):return pd.Timestamp(c).year,float(v)
    return None

def dcf_inputs(frames):
    cf=frames.get('CashFlow'); bs=frames.get('Balance')
    fcf=_latest(cf,'Free Cash Flow'); fcf_method='provider Free Cash Flow'
    if fcf is None:
        cfo=_latest(cf,'Operating Cash Flow'); capex=_latest(cf,'Capital Expenditure')
        if cfo and capex:
            capex_outflow=abs(capex[1]); fcf=(cfo[0],cfo[1]-capex_outflow); fcf_method='Operating Cash Flow - |Capital Expenditure|'
    debt=_latest(bs,'Total Debt'); cash=_latest(bs,'Cash And Cash Equivalents')
    missing=[]
    if fcf is None:missing.append('free cash flow (or operating cash flow + capex)')
    if debt is None:missing.append('total debt')
    if cash is None:missing.append('cash and cash equivalents')
    return {'fcf':fcf,'fcf_method':fcf_method,'debt':debt,'cash':cash,'missing':missing}

def validate(a):
    if not 3<=a.n<=10:raise ValueError('Forecast years must be 3–10')
    if not -.25<=a.growth<=.40:raise ValueError('FCF growth assumption is outside the supported learning range (-25% to 40%).')
    if not .03<=a.wacc<=.30:raise ValueError('WACC must be between 3% and 30%.')
    if not -.02<=a.terminal_growth<=.08:raise ValueError('Terminal growth must be between -2% and 8%.')
    if a.wacc<=a.terminal_growth:raise ValueError('WACC must be greater than terminal growth.')
    if a.shares is not None and a.shares<=0:raise ValueError('Shares outstanding must be positive.')

def value_company(frames,a:Assumptions):
    validate(a);x=dcf_inputs(frames)
    if x['missing']:raise ValueError('Valuation unavailable: missing '+', '.join(x['missing']))
    year,fcf=x['fcf'];flows=[];cur=fcf
    for i in range(1,a.n+1):
        cur*=1+a.growth;flows.append({'year':year+i,'fcf_raw':cur,'discount_factor':1/(1+a.wacc)**i,'pv_raw':cur/(1+a.wacc)**i})
    terminal=cur*(1+a.terminal_growth)/(a.wacc-a.terminal_growth)
    terminal_pv=terminal/(1+a.wacc)**a.n
    ev=sum(r['pv_raw'] for r in flows)+terminal_pv
    equity=ev-x['debt'][1]+x['cash'][1]
    per_share=equity/a.shares if a.shares and a.shares>0 else None
    diff=per_share/a.market_price-1 if per_share is not None and a.market_price and a.market_price>0 else None
    terminal_share=terminal_pv/ev if ev else None
    return {'base_year':year,'base_fcf_raw':fcf,'fcf_method':x['fcf_method'],'flows':flows,'terminal_value_raw':terminal,'terminal_pv_raw':terminal_pv,
            'terminal_value_share':terminal_share,'enterprise_value_raw':ev,'equity_value_raw':equity,
            'intrinsic_value_per_share':per_share,'difference_vs_price':diff,'debt_raw':x['debt'][1],'cash_raw':x['cash'][1]}

def scenarios(frames,a):
    configs={
        'Conservative':replace(a,growth=max(-.25,a.growth-.02),wacc=min(.30,a.wacc+.01),terminal_growth=max(-.02,a.terminal_growth-.005)),
        'Base':a,
        'Optimistic':replace(a,growth=min(.40,a.growth+.02),wacc=max(.03,a.wacc-.01),terminal_growth=min(.08,a.terminal_growth+.005)),
    }
    out=[]
    for name,ass in configs.items():
        if ass.wacc<=ass.terminal_growth:continue
        v=value_company(frames,ass);out.append({'scenario':name,'growth':ass.growth,'wacc':ass.wacc,'terminal_growth':ass.terminal_growth,
            'enterprise_value_raw':v['enterprise_value_raw'],'equity_value_raw':v['equity_value_raw'],'intrinsic_value_per_share':v['intrinsic_value_per_share']})
    return out

def sensitivity(frames,a):
    rows=[]
    for dw in (-.01,0,.01):
        for dg in (-.005,0,.005):
            aa=replace(a,wacc=a.wacc+dw,terminal_growth=a.terminal_growth+dg)
            if aa.wacc<=aa.terminal_growth:continue
            v=value_company(frames,aa);rows.append({'wacc':aa.wacc,'terminal_growth':aa.terminal_growth,'intrinsic_value_per_share':v['intrinsic_value_per_share'],'equity_value_raw':v['equity_value_raw']})
    return rows

def workbook(ticker,frames,a):
    v=value_company(frames,a);sc=scenarios(frames,a);sens=sensitivity(frames,a)
    wb=Workbook();ws=wb.active;ws.title='Valuation'
    ws.append(['FinSight AI — Valuation Playground',ticker]);ws.append(['Purpose','Educational scenario, not a company forecast or investment recommendation']);ws.append([])
    ws.append(['Input','Value']);ws.append(['Base FCF (raw provider units)',v['base_fcf_raw']]);ws.append(['FCF source/method',v['fcf_method']]);ws.append(['FCF growth',a.growth]);ws.append(['WACC',a.wacc]);ws.append(['Terminal growth',a.terminal_growth]);ws.append(['Forecast years',a.n]);ws.append(['Shares outstanding',a.shares]);ws.append(['Market price / share',a.market_price])
    ws.append([]);ws.append(['Output','Value']);ws.append(['Enterprise value (raw provider units)',v['enterprise_value_raw']]);ws.append(['Equity value (raw provider units)',v['equity_value_raw']]);ws.append(['Illustrative intrinsic value / share',v['intrinsic_value_per_share']]);ws.append(['Difference vs entered market price',v['difference_vs_price']]);ws.append(['PV of terminal value / enterprise value',v['terminal_value_share']])
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='7657D6')
    f=wb.create_sheet('FCF Forecast');f.append(['Year','FCF raw units','Discount factor','PV raw units']);[f.append([r['year'],r['fcf_raw'],r['discount_factor'],r['pv_raw']]) for r in v['flows']]
    s=wb.create_sheet('Scenarios');s.append(['Scenario','FCF growth','WACC','Terminal growth','Enterprise value raw','Equity value raw','Value/share']);[s.append([r[k] for k in ['scenario','growth','wacc','terminal_growth','enterprise_value_raw','equity_value_raw','intrinsic_value_per_share']]) for r in sc]
    z=wb.create_sheet('Sensitivity');z.append(['WACC','Terminal growth','Value/share','Equity value raw']);[z.append([r['wacc'],r['terminal_growth'],r['intrinsic_value_per_share'],r['equity_value_raw']]) for r in sens]
    n=wb.create_sheet('Source Notes');n.append(['Provider','Yahoo Finance / yfinance']);n.append(['Status','Provider extraction; verify against company filings']);n.append(['Model','Simple FCFF growth scenario; no automatic three-statement projection']);n.append(['Critical note','Share price and shares are user/provider inputs and should be independently verified.'])
    out=io.BytesIO();wb.save(out);return out.getvalue(),v,sc,sens
